"""Render a shared resume layout as an editable Word or LaTeX document."""
from io import BytesIO


def _items(value):
    return value if isinstance(value, list) else []


def _sections(resume):
    """Shared section order and content source for both export formats."""
    sections = (
        ("summary", "Professional Summary", resume.get("summary")),
        ("skills", "Skills", _items(resume.get("skills"))),
        ("experience", "Experience", _items(resume.get("experience"))),
        ("projects", "Projects", _items(resume.get("projects"))),
        ("education", "Education", _items(resume.get("education"))),
        ("certifications", "Certifications", _items(resume.get("certifications"))),
    )
    return [(kind, title, content) for kind, title, content in sections if content]


def _style(resume):
    style = resume.get("style") if isinstance(resume.get("style"), dict) else {}

    def valid(value, default):
        if isinstance(value, str) and len(value) == 6 and all(char in "0123456789abcdefABCDEF" for char in value):
            return value.upper()
        return default

    return {
        "accent": valid(style.get("accent_color"), "222222"),
        "secondary": valid(style.get("secondary_color"), "414141"),
    }


def _project_dict(project):
    if isinstance(project, str):
        return {"name": project}
    return project if isinstance(project, dict) else None


def build_docx(resume: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    page = doc.sections[0]
    page.top_margin = Inches(0.55)
    page.bottom_margin = Inches(0.55)
    page.left_margin = Inches(0.7)
    page.right_margin = Inches(0.7)

    palette = _style(resume)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    title_style = doc.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(24)
    title_style.font.color.rgb = RGBColor.from_string(palette["accent"])
    for style_name, color, size in (("Heading 1", palette["accent"], 11), ("Heading 2", palette["secondary"], 10)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Heading 1" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 5)
        style.paragraph_format.space_after = Pt(4 if style_name == "Heading 1" else 2)

    doc.add_heading(str(resume.get("name") or "Resume"), 0)
    contact = " | ".join(
        str(resume.get(key, "")).strip()
        for key in ("email", "phone", "location", "linkedin", "website")
        if resume.get(key)
    )
    if contact:
        doc.add_paragraph(contact)

    for kind, title, content in _sections(resume):
        doc.add_heading(title, level=1)
        if kind == "summary":
            doc.add_paragraph(str(content))
        elif kind == "skills":
            doc.add_paragraph(" • ".join(map(str, content)))
        elif kind == "experience":
            for job in content:
                heading = " — ".join(str(job.get(key, "")).strip() for key in ("title", "company") if job.get(key))
                if job.get("duration"):
                    heading += f" | {job['duration']}"
                doc.add_heading(heading or "Experience", level=2)
                for bullet in _items(job.get("bullets")):
                    doc.add_paragraph(str(bullet), style="List Bullet")
        elif kind == "projects":
            for raw_project in content:
                project = _project_dict(raw_project)
                if not project:
                    continue
                heading = str(project.get("name") or "Project").strip()
                if project.get("duration"):
                    heading += f" | {project['duration']}"
                doc.add_heading(heading, level=2)
                technologies = _items(project.get("technologies"))
                if technologies:
                    doc.add_paragraph("Technologies: " + ", ".join(map(str, technologies)))
                if project.get("description"):
                    doc.add_paragraph(str(project["description"]))
                for bullet in _items(project.get("bullets")):
                    doc.add_paragraph(str(bullet), style="List Bullet")
                if project.get("url"):
                    doc.add_paragraph("Project link: " + str(project["url"]))
        elif kind == "education":
            for item in content:
                text = " — ".join(str(item.get(key, "")).strip() for key in ("degree", "institution", "year") if item.get(key))
                if item.get("gpa"):
                    text += f" | GPA: {item['gpa']}"
                doc.add_paragraph(text, style="List Bullet")
        elif kind == "certifications":
            for item in content:
                doc.add_paragraph(str(item), style="List Bullet")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _tex(value) -> str:
    replacements = {
        "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
        "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
        "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(char, char) for char in str(value or ""))


def build_latex(resume: dict) -> str:
    palette = _style(resume)
    lines = [
        r"\documentclass[10pt]{article}",
        r"\usepackage[margin=0.65in]{geometry}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{textcomp}",
        r"\usepackage{enumitem}",
        r"\usepackage{xcolor}",
        r"\usepackage[hidelinks]{hyperref}",
        r"\setlist[itemize]{nosep,leftmargin=*}",
        r"\pagestyle{empty}",
        r"\definecolor{ResumeAccent}{HTML}{" + palette["accent"] + "}",
        r"\definecolor{ResumeSecondary}{HTML}{" + palette["secondary"] + "}",
        r"\newcommand{\resumeSection}[1]{\section*{\color{ResumeAccent}#1}}",
        r"\begin{document}",
        r"\begin{center}",
        r"{\LARGE \textcolor{ResumeAccent}{\textbf{" + _tex(resume.get("name") or "Resume") + r"}}}\\[4pt]",
    ]
    contact = [resume.get(key) for key in ("email", "phone", "location", "linkedin", "website") if resume.get(key)]
    if contact:
        lines.append(r" \textbar{} ".join(_tex(item) for item in contact) + r"\\")
    lines.extend([r"\end{center}", r"\vspace{-6pt}"])

    for kind, title, content in _sections(resume):
        lines.append(r"\resumeSection{" + _tex(title) + "}")
        if kind == "summary":
            lines.append(_tex(content))
        elif kind == "skills":
            lines.append(r" \textbar{} ".join(_tex(skill) for skill in content))
        elif kind == "experience":
            for job in content:
                heading = " -- ".join(str(job.get(key, "")).strip() for key in ("title", "company") if job.get(key))
                if job.get("duration"):
                    heading += " | " + str(job["duration"])
                lines.append(r"\textcolor{ResumeSecondary}{\textbf{" + _tex(heading) + r"}}\\")
                bullets = _items(job.get("bullets"))
                if bullets:
                    lines.append(r"\begin{itemize}")
                    lines.extend(r"\item " + _tex(item) for item in bullets)
                    lines.append(r"\end{itemize}")
        elif kind == "projects":
            for raw_project in content:
                project = _project_dict(raw_project)
                if not project:
                    continue
                heading = str(project.get("name") or "Project").strip()
                if project.get("duration"):
                    heading += " | " + str(project["duration"])
                lines.append(r"\textcolor{ResumeSecondary}{\textbf{" + _tex(heading) + r"}}\\")
                technologies = _items(project.get("technologies"))
                if technologies:
                    lines.append(r"\textit{Technologies:} " + ", ".join(_tex(item) for item in technologies) + r"\\")
                if project.get("description"):
                    lines.append(_tex(project["description"]) + r"\\")
                if project.get("url"):
                    lines.append(r"\textit{Project link:} " + _tex(project["url"]) + r"\\")
                bullets = _items(project.get("bullets"))
                if bullets:
                    lines.append(r"\begin{itemize}")
                    lines.extend(r"\item " + _tex(item) for item in bullets)
                    lines.append(r"\end{itemize}")
        elif kind == "education":
            lines.append(r"\begin{itemize}")
            for item in content:
                value = " -- ".join(str(item.get(key, "")).strip() for key in ("degree", "institution", "year") if item.get(key))
                if item.get("gpa"):
                    value += " | GPA: " + str(item["gpa"])
                lines.append(r"\item " + _tex(value))
            lines.append(r"\end{itemize}")
        elif kind == "certifications":
            lines.append(r"\begin{itemize}")
            lines.extend(r"\item " + _tex(item) for item in content)
            lines.append(r"\end{itemize}")
    lines.append(r"\end{document}")
    return "\n".join(lines) + "\n"
